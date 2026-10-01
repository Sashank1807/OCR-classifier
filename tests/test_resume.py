import pytest
import docx
import pathlib
from app.services.word_service import word_service
from app.services.ocr_pipeline import pipeline
from app.core.database import SessionLocal
from app.utils.structured_schemas import normalize_structured_data


def test_word_resume_extraction(tmp_path):
    doc_path = tmp_path / "sample_resume.docx"
    doc = docx.Document()
    doc.add_heading("Jane Doe", level=1)
    doc.add_paragraph("Senior Python Developer | Email: jane.doe@example.com | Phone: +1 555-0199")
    doc.add_heading("Summary", level=2)
    doc.add_paragraph("Over 8 years of experience developing enterprise AI platforms and REST APIs.")
    doc.add_heading("Technical Skills", level=2)
    doc.add_paragraph("Python, FastAPI, PyTorch, SQL, Docker, Kubernetes, AWS")
    doc.add_heading("Work Experience", level=2)
    doc.add_paragraph("Lead AI Engineer - TechCorp Solutions (2020 - Present)")
    doc.add_paragraph("Architected scalable OCR intelligence microservices.")
    doc.add_heading("Education", level=2)
    doc.add_paragraph("B.S. in Computer Science - State University (2016)")
    doc.save(str(doc_path))

    md, plain, pages = word_service.extract_word_content(doc_path)
    assert "Jane Doe" in md
    assert "jane.doe@example.com" in md
    assert "Technical Skills" in md
    assert len(pages) == 1


def test_pipeline_resume_docx(tmp_path):
    doc_path = tmp_path / "jane_resume.docx"
    doc = docx.Document()
    doc.add_heading("Jane Doe - Resume", level=1)
    doc.add_paragraph("Email: jane@test.com | Phone: 9876543210")
    doc.add_heading("Education", level=2)
    doc.add_paragraph("B.Tech in Computer Science")
    doc.add_heading("Skills", level=2)
    doc.add_paragraph("Python, OCR, Machine Learning")
    doc.save(str(doc_path))

    db = SessionLocal()
    try:
        res = pipeline.process_file(db, doc_path, "jane_resume.docx", requested_doc_type="RESUME")
        assert res["document_type"] == "RESUME"
        assert "Jane Doe" in res["markdown"]
        # The API publishes markdown, not plain_text - the two carried the same
        # content and plain_text was dropped from the payload.
        assert "jane@test.com" in res["markdown"]
        assert "plain_text" not in res
        assert res["status"] == "COMPLETED"
    finally:
        db.close()


def test_resume_structured_schema():
    raw = {
        "candidate_name": "Alex Johnson",
        "contact_info": {
            "email": "alex@example.com",
            "phone": "+1 234 567 8900",
            "linkedin": "linkedin.com/in/alex"
        },
        "professional_summary": "Full stack engineer.",
        "skills": {
            "technical_skills": ["Python", "React", "TypeScript"]
        },
        "education": [
            {"degree": "B.S. Computer Science", "institution": "MIT"}
        ]
    }
    norm = normalize_structured_data("RESUME", raw)
    assert norm["document_type"] == "RESUME"
    assert norm["fields"]["candidate_name"] == "Alex Johnson"
    assert norm["fields"]["contact_info"]["email"] == "alex@example.com"
    assert norm["overall_confidence"] >= 0.80
